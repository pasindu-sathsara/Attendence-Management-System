import csv
import io
import json
import tempfile
import unittest
import zipfile
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage
import app as backend
import test_reports


class StudentReportTests(test_reports.ReportTests):
    def download(self, registrations):
        with patch.object(backend, 'query', side_effect=self.query):
            return self.client.post('/api/student-reports.zip?academic_year=26/27&type=scholarship',
                json={'registrations': registrations}, headers={'X-CSRF-Token': 'token'})

    def test_one_selected_and_all_students_get_separate_files(self):
        for selected in [['ICT/2024/001'], ['ICT/2024/002', 'ICT/2025/001'],
                         ['ICT/2024/001', 'ICT/2024/002', 'ICT/2025/001']]:
            response = self.download(selected)
            self.assertEqual(response.status_code, 200)
            with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
                self.assertEqual(len(archive.namelist()), len(selected))
                found = set()
                for filename in archive.namelist():
                    book = load_workbook(io.BytesIO(archive.read(filename)))
                    summary = list(book['Student Summary'].values)
                    reg = summary[1][0]
                    found.add(reg)
                    self.assertEqual(len(summary), 2)
                    self.assertEqual(summary[0][-1], 'Eligibility')
                    for row in list(book['Session Attendance'].values)[1:]:
                        self.assertIn(row[0], ['ICT 1305(T)', 'COM 1301(T)'])
                    expected = 'Eligible' if reg == 'ICT/2024/002' else 'Ineligible'
                    self.assertEqual(summary[1][-1], expected)
                self.assertEqual(found, set(selected))

    def test_more_than_120_has_no_fixed_student_limit(self):
        regs = []
        for number in range(121):
            reg = f'ICT/2026/{number:03}'
            email = 'new' + str(number)
            regs.append(reg)
            self.conn.execute('INSERT INTO student VALUES (?,?,?)', (email, reg, 'Student ' + str(number)))
            self.conn.execute("INSERT INTO student_course VALUES (?,'COM 1301','26/27')", (email,))
        response = self.download(regs)
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            self.assertEqual(len(archive.namelist()), 121)
            book = load_workbook(io.BytesIO(archive.read(archive.namelist()[0])))
            self.assertEqual(list(book.active.values)[1][-1], 'Not assessed')

    def test_invalid_empty_and_out_of_scope_selection_rejected(self):
        for regs in [[], ['unknown'], [42], ['ICT/2024/001', 'unknown']]:
            self.assertEqual(self.download(regs).status_code, 400)
        for role in ['lecturer', 'student']:
            with self.client.session_transaction() as state:
                state['role'] = role
            self.assertEqual(self.download(['ICT/2024/001']).status_code, 403)

    def test_scholarship_summary_and_excel_include_eligibility(self):
        response = self.get('/api/report?academic_year=26/27&type=scholarship')
        self.assertEqual([r['eligibility'] for r in response.json['rows']],
                         ['Ineligible', 'Eligible', 'Ineligible'])
        response = self.get('/api/report.xlsx?academic_year=26/27&type=scholarship')
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual(list(book.active.values)[0][-1], 'Eligibility')
        self.assertEqual(list(book.active.values)[1][-1], 'Ineligible')


class SetupAndImportTests(unittest.TestCase):
    def setUp(self):
        self.client = backend.app.test_client()
        with self.client.session_transaction() as state:
            state.update(role='hod', email='h@example.com', csrf='token')

    def post(self, path, body):
        return self.client.post(path, json=body, headers={'X-CSRF-Token': 'token'})

    def test_semester_dates_are_saved_and_can_be_updated(self):
        body = dict(academic_year='26/27', study_year=1, semester_number=1,
                    start_date='2026-07-13', end_date='2026-10-23')
        with patch.object(backend, 'query', return_value={'email': 'h@example.com'}), \
             patch.object(backend, 'execute') as execute, patch.object(backend, 'db'):
            self.assertEqual(self.post('/api/semesters', body).status_code, 200)
            self.assertIn('ON DUPLICATE KEY UPDATE', execute.call_args.args[0])
            self.assertEqual(execute.call_args.args[1][-2:], (date(2026, 7, 13), date(2026, 10, 23)))
            execute.reset_mock()
            body['end_date'] = '2026-01-01'
            self.assertEqual(self.post('/api/semesters', body).status_code, 400)
            execute.assert_not_called()

    def test_all_components_use_same_semester_end_for_special_medical(self):
        for code in ['ICT 1407', 'ICT 1407(P)', 'COM 1301']:
            course = dict(course_code=code, academic_year='26/27', study_year=1, semester_number=1)
            with patch.object(backend, 'query', return_value={'end_date': date.today()-timedelta(days=1)}) as query:
                self.assertTrue(backend.medical_rules(course, date.today()-timedelta(days=30))['special'])
                self.assertEqual(query.call_args.args[1], ('26/27', 1, 1))
            with patch.object(backend, 'query', return_value={'end_date': None}):
                self.assertFalse(backend.medical_rules(course, date.today()-timedelta(days=30))['special'])

    def test_template_contains_existing_and_new_students_with_blank_account_fields(self):
        with tempfile.TemporaryDirectory() as folder:
            stage = Path(folder)/'.imports'
            stage.mkdir()
            (stage/'token').write_text(json.dumps({'rows': [dict(reg='ICT/2024/001', email='old@example.com'), dict(reg='ICT/2024/002', email='')]}))
            with self.client.session_transaction() as state:
                state['import_token'] = 'token'
            with patch.object(backend, 'ROOT', Path(folder)), patch.object(backend, 'query', return_value={'email':'h@example.com'}):
                response = self.client.get('/api/account-template')
            rows = list(csv.reader(io.StringIO(response.text)))
            self.assertEqual(rows, [['Reg_no','Email','NIC'], ['ICT/2024/001','',''], ['ICT/2024/002','','']])

    def test_existing_accounts_skip_password_hashing(self):
        file = FileStorage(stream=io.BytesIO(b'Reg_no,Email,NIC\nICT/2024/001,old@example.com,200012345678\nICT/2024/002,new@example.com,200012345679\n'), filename='accounts.csv')
        with patch.object(backend, 'generate_password_hash', return_value='secure-hash') as hashing:
            rows = backend.account_details(file, {'ICT/2024/001': {'email':'old@example.com'}})
            hashing.assert_called_once_with('200012345679')
            self.assertEqual(rows['ICT/2024/001']['password_hash'], '')
            self.assertNotIn('nic', rows['ICT/2024/002'])

    def test_preview_uses_bulk_lookup_and_never_stages_raw_nic(self):
        modules = [dict(code='ICT 1305', study=1, sheet='ICT 1305', students=[dict(reg='ICT/2024/001', name='A'), dict(reg='ICT/2024/002',name='B')])]
        with tempfile.TemporaryDirectory() as folder, patch.object(backend, 'ROOT', Path(folder)), \
             patch.object(backend, 'query', side_effect=[{'email':'h@example.com'}, [{'reg_number':'ICT/2024/001','email':'old@example.com'}]]) as query, \
             patch.object(backend, 'workbook_rows', return_value=modules):
            response = self.client.post('/api/import-preview', data={'semester':'26/27|1|1', 'file':(io.BytesIO(b'x'),'file.xlsx'),
                'accounts':(io.BytesIO(b'Reg_no,Email,NIC\nICT/2024/002,new@example.com,200012345679\n'),'accounts.csv')}, headers={'X-CSRF-Token':'token'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(query.call_count, 2)
            self.assertIn('IN (%s,%s)', query.call_args.args[0])
            staged = next((Path(folder)/'.imports').iterdir()).read_text()
            self.assertNotIn('200012345679', staged)
            self.assertEqual(response.json['missing'], 0)


if __name__ == '__main__':
    unittest.main()
