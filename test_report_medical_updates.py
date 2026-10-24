"""Report selection, medical deadlines and review permissions using isolated test data."""
import io
import sqlite3
import unittest
from datetime import date, datetime, timedelta
from unittest.mock import patch
from openpyxl import load_workbook
import app as backend
from test_reports import ReportTests
from test_new_features import NewFeatureTests


class ReportSelectionTests(ReportTests):
    def post(self, path, registrations):
        with patch.object(backend, 'query', side_effect=self.query):
            return self.client.post(path, json={'registrations': registrations}, headers={'X-CSRF-Token':'token'})

    def test_selected_students_apply_to_summary_registers_and_excel(self):
        path='?academic_year=26/27&batch=23/24&type=full'
        response=self.post('/api/report'+path,['ICT/2024/001'])
        module=next(r for r in response.json['rows'] if r['course_code']=='ICT 1305')
        self.assertEqual(module['students'],1)
        self.assertEqual(module['percentage'],80)
        detail=self.post('/api/report-view'+path,['ICT/2024/001']).json['registers']
        for register in detail:
            self.assertEqual({r['reg'] for r in register['rows']},{'ICT/2024/001'})
        response=self.post('/api/report.xlsx'+path,['ICT/2024/001'])
        book=load_workbook(io.BytesIO(response.data))
        self.assertEqual(set(book.sheetnames),{'COM 1301(T)','ICT 1305(T)'})
        records=[row for sheet in book for row in sheet.values if isinstance(row[0],int)]
        self.assertEqual({row[1] for row in records},{'ICT/2024/001'})
        self.assertEqual(len(records),2)

    def test_theory_and_practical_have_distinct_tabs_and_exact_rosters(self):
        self.conn.execute("INSERT INTO course VALUES ('ICT 1305(P)',1,1)")
        self.conn.execute("INSERT INTO course_offering VALUES ('ICT 1305(P)','26/27',NULL)")
        self.conn.execute("INSERT INTO student_course VALUES ('b','ICT 1305(P)','26/27')")
        response=self.get('/api/report.xlsx?academic_year=26/27&batch=23/24')
        book=load_workbook(io.BytesIO(response.data))
        self.assertEqual(set(book.sheetnames),{'COM 1301(T)','ICT 1305(T)','ICT 1305(P)'})
        theory=[row[1] for row in book['ICT 1305(T)'].values if isinstance(row[0],int)]
        practical=[row[1] for row in book['ICT 1305(P)'].values if isinstance(row[0],int)]
        self.assertEqual(theory,['ICT/2024/001','ICT/2024/002'])
        self.assertEqual(practical,['ICT/2024/002'])
        self.assertEqual(book['ICT 1305(P)'].freeze_panes,'D8')

    def test_all_approved_reasons_count_in_reports_export_and_prediction(self):
        for reason,field,code in [('MC','medical','MC'),('SMC','special','SMC'),('EXCUSE','excused','EX')]:
            self.conn.execute("UPDATE student_class_session SET absent_reason=? WHERE email='a' AND week_number=5",(reason,))
            response=self.get('/api/report?course=ICT+1305~26/27&batch=23/24')
            row=response.json['rows'][0]
            self.assertEqual((row['present'],row['absent'],row[field],row['percentage'],row['eligibility']),(4,0,1,100,'Eligible'))
            single=self.get('/api/report?academic_year=26/27&type=single&batch=23/24').json['rows'][0]
            self.assertEqual(single['percentage'],100)
            self.assertEqual(single['eligibility'],'Eligible')
            response=self.get('/api/report.xlsx?course=ICT+1305~26/27&batch=23/24')
            record=list(load_workbook(io.BytesIO(response.data)).active.values)[7]
            self.assertEqual(record[7],code)
            self.assertEqual(record[-2:],(100,'Eligible'))
        with self.client.session_transaction() as state:state.update(role='student',email='a')
        row=next(r for r in self.get('/api/my-attendance').json['rows'] if r['course_code']=='ICT 1305')
        self.assertEqual(row['credited'],5)
        self.assertEqual(row['maximum'],100)
        self.assertEqual(row['prediction'],'Above threshold')

    def test_filtered_student_picker_and_invalid_selections(self):
        response=self.get('/api/report-students?academic_year=26/27&batch=24/25')
        self.assertEqual([r['reg'] for r in response.json['rows']],['ICT/2025/001'])
        for registrations in [[],['unknown'],['ICT/2025/001'],['ICT/2024/001','unknown'],[42]]:
            for endpoint in ['/api/report','/api/report-view','/api/report.xlsx']:
                self.assertEqual(self.post(endpoint+'?academic_year=26/27&batch=23/24',registrations).status_code,400)

    def test_removed_scholarship_type_is_not_exposed(self):
        self.assertEqual(self.get('/api/report?academic_year=26/27&type=scholarship').status_code,400)
        with patch.object(backend,'query',side_effect=self.query):
            page=self.client.get('/reports.html')
        self.assertNotIn(b'id="single-report"',page.data)
        self.assertNotIn(b'id="full-report"',page.data)
        self.assertIn(b'Download Separate Student Reports (ZIP)',page.data)
        self.assertNotIn(b'Mahapola',page.data)
        self.assertNotIn(b'report-course',page.data)

    def test_excused_counts_towards_attendance(self):
        self.conn.execute("UPDATE student_class_session SET absent_reason='EXCUSE' WHERE email='a' AND week_number=5")
        response=self.get('/api/report?course=ICT+1305~26/27&batch=23/24')
        row=response.json['rows'][0]
        self.assertEqual((row['present'],row['absent'],row['excused'],row['percentage']),(4,0,1,100))
        response=self.get('/api/report-view?course=ICT+1305~26/27&batch=23/24')
        self.assertEqual(response.json['registers'][0]['rows'][0]['marks'][-1],'EX')

    def test_unmarked_is_a_dash_in_preview_and_export(self):
        self.conn.execute("DELETE FROM student_class_session WHERE email='a' AND week_number=5")
        response=self.get('/api/report-view?course=ICT+1305~26/27&batch=23/24')
        self.assertEqual(response.json['registers'][0]['rows'][0]['marks'][-1],'—')
        response=self.get('/api/report.xlsx?course=ICT+1305~26/27&batch=23/24')
        self.assertEqual(list(load_workbook(io.BytesIO(response.data)).active.values)[7][7],'—')

    def test_lecturer_assigned_module_only_and_all_enrolled_students(self):
        self.conn.executescript("INSERT INTO lecturer VALUES ('h@example.com','Teacher','Mr'); INSERT INTO lecturer_course VALUES ('h@example.com','ICT 1305','26/27');")
        with self.client.session_transaction() as state:state['role']='lecturer'
        def query(sql,args=(),one=False):
            if sql.startswith('SELECT email FROM lecturer_course'):
                return {'email':'h@example.com'} if args[1]=='ICT 1305' else None
            return self.query(sql,args,one)
        with patch.object(backend,'query',side_effect=query):
            self.assertEqual(len(self.client.get('/api/report?course=ICT+1305~26/27').json['rows']),3)
            self.assertEqual(self.client.get('/api/report?course=COM+1301~26/27').status_code,403)
            self.assertEqual(self.client.get('/api/report?academic_year=26/27').status_code,403)
            self.assertEqual(self.client.get('/api/report-students?academic_year=26/27').status_code,403)
            self.assertEqual(self.client.post('/api/report?course=ICT+1305~26/27',json={'registrations':['ICT/2024/001']},headers={'X-CSRF-Token':'token'}).status_code,403)


class ExcusedUpdateTests(NewFeatureTests):
    def test_excused_is_saved_in_existing_absent_reason(self):
        response,calls=self.medical_save('lecturer','excused',False)
        self.assertEqual(response.status_code,200)
        self.assertEqual(calls[0].args[1][-2:],(0,'EXCUSE'))
        self.assertEqual(self.medical_save('lecturer','excused',False,current_status=1)[0].status_code,400)
        self.assertEqual(self.medical_save('ar','excused',False)[0].status_code,403)


class MedicalSubmissionTests(unittest.TestCase):
    def setUp(self):
        self.client=backend.app.test_client()
        self.signin('student','a')
        sqlite3.register_converter('DATE',lambda b: date.fromisoformat(b.decode()))
        sqlite3.register_converter('TIMESTAMP',lambda b:datetime.fromisoformat(b.decode()))
        self.conn=sqlite3.connect(':memory:',detect_types=sqlite3.PARSE_DECLTYPES)
        self.conn.row_factory=sqlite3.Row
        self.conn.create_function('NOW',0,lambda:datetime.now().isoformat(' '))
        from test_options import sql_functions
        sql_functions(self.conn)
        self.conn.executescript('''
            CREATE TABLE student(email TEXT,reg_number TEXT,name TEXT);
            CREATE TABLE student_course(email TEXT,course_code TEXT,academic_year TEXT);
            CREATE TABLE lecturer_course(email TEXT,course_code TEXT,academic_year TEXT);
            CREATE TABLE course(course_code TEXT,study_year INT,semester_number INT);
            CREATE TABLE course_offering(course_code TEXT,academic_year TEXT,teaching_end DATE);
            CREATE TABLE semester(academic_year TEXT,study_year INT,semester_number INT,end_date DATE);
            CREATE TABLE class_session(course_code TEXT,academic_year TEXT,date DATE,week_number INT,updated_by TEXT,updated_time TIMESTAMP,email TEXT);
            CREATE TABLE student_class_session(email TEXT,course_code TEXT,date DATE,week_number INT,
                attendance_status INT,absent_reason TEXT,medical_report BLOB,medical_filename TEXT,
                medical_mimetype TEXT,medical_kind TEXT,medical_status TEXT,medical_submitted_at TIMESTAMP);
            INSERT INTO student VALUES ('a','ICT/2024/001','Alpha'),('b','ICT/2024/002','Beta');
            INSERT INTO course VALUES ('ICT1305',1,1);
            INSERT INTO course_offering VALUES ('ICT1305','26/27',NULL);
            INSERT INTO student_course VALUES ('a','ICT1305','26/27'),('b','ICT1305','26/27');
            INSERT INTO lecturer_course VALUES ('teacher','ICT1305','26/27');
        ''')
        self.when=date.today()-timedelta(days=5)
        self.conn.execute('INSERT INTO semester VALUES (?,?,?,?)',('26/27',1,1,date.today()-timedelta(days=1)))
        self.conn.execute('INSERT INTO class_session(course_code,academic_year,date,week_number) VALUES (?,?,?,?)',('ICT1305','26/27',self.when,1))
        for email in ['a','b']:
            self.conn.execute('INSERT INTO student_class_session(email,course_code,date,week_number,attendance_status) VALUES (?,?,?,?,?)',(email,'ICT1305',self.when,1,0))
        self.patches=[patch.object(backend,'query',side_effect=self.query),patch.object(backend,'execute',side_effect=self.execute),patch.object(backend,'db')]
        for item in self.patches:item.start()

    def tearDown(self):
        for item in reversed(self.patches):item.stop()
        self.conn.close()

    def signin(self,role,email):
        with self.client.session_transaction() as state:state.update(role=role,email=email,name='Mr.Teacher',csrf='token')

    def query(self,sql,args=(),one=False):
        if sql.startswith('SELECT email,must_change_password FROM'):
            return {'email':'account','must_change_password':0}
        cursor=self.conn.execute(sql.replace(' FOR UPDATE','').replace('RIGHT(', 'RIGHT_STR(').replace('%s','?'),args)
        if one:
            row=cursor.fetchone();return dict(row) if row else None
        return [dict(r) for r in cursor.fetchall()]

    def execute(self,sql,args=()):
        return self.conn.execute(sql.replace('%s','?'),args).rowcount

    def key(self):return dict(course='ICT1305~26/27',session=str(self.when)+'|1')

    def upload(self,kind='normal',content=b'%PDF-1.4\nmedical'):
        return self.client.post('/api/medical-upload',data={**self.key(),'kind':kind,'email':'b',
            'file':(io.BytesIO(content),'medical.pdf')},headers={'X-CSRF-Token':'token'})

    def review(self,decision='approved'):
        return self.client.post('/api/medical-review',json={**self.key(),'email':'a','decision':decision},headers={'X-CSRF-Token':'token'})

    def record(self,email='a'):
        return dict(self.conn.execute('SELECT * FROM student_class_session WHERE email=?',(email,)).fetchone())

    def test_upload_stays_pending_and_is_bound_to_signed_in_student(self):
        self.assertEqual(self.upload().status_code,200)
        row=self.record();self.assertEqual(row['medical_status'],'pending')
        self.assertEqual(row['attendance_status'],0);self.assertIsNone(row['absent_reason'])
        self.assertIsNone(self.record('b')['medical_report'])
        self.assertEqual(self.upload().status_code,409)
        history=self.client.get('/api/my-sessions?course=ICT1305~26/27').json['rows'][0]
        self.assertEqual(history['medical_status'],'pending');self.assertFalse(history['normal'])

    def test_excuse_letter_approval_changes_status_and_percentage(self):
        self.assertEqual(self.upload('excuse').status_code,200)
        before=self.client.get('/api/my-attendance').json['rows'][0]
        self.assertEqual((before['percentage'],before['absent']),(0,1))
        self.signin('lecturer','teacher')
        listing=self.client.get('/api/medical-submissions').json
        self.assertEqual(listing['rows'][0]['medical_kind'],'excuse')
        self.assertEqual(self.review().status_code,200)
        row=self.record()
        self.assertEqual((row['attendance_status'],row['absent_reason'],row['medical_status']),(0,'EXCUSE','approved'))
        self.signin('student','a')
        after=self.client.get('/api/my-attendance').json['rows'][0]
        self.assertEqual((after['present'],after['excused'],after['absent'],after['percentage']),(0,1,0,100))
        history=self.client.get('/api/my-sessions?course=ICT1305~26/27').json['rows'][0]
        self.assertEqual(history['status'],'excused')
        self.assertEqual(history['medical_kind'],'excuse')

    def test_excuse_rejection_and_access_deadline_rules(self):
        self.assertEqual(self.upload('excuse').status_code,200)
        for role,email in [('lecturer','other'),('ar','ar'),('hod','head')]:
            self.signin(role,email)
            self.assertEqual(self.client.get('/api/medical-submissions').json['total'],0)
            self.assertEqual(self.review().status_code,403)
        self.signin('lecturer','teacher');self.assertEqual(self.review('rejected').status_code,200)
        self.assertIsNone(self.record()['absent_reason'])
        self.signin('student','a')
        self.when=date.today()-timedelta(days=15)
        self.conn.execute('UPDATE class_session SET date=?',(self.when,))
        self.conn.execute('UPDATE student_class_session SET date=?',(self.when,))
        self.assertEqual(self.upload('excuse').status_code,400)

    def test_normal_approval_updates_absence_and_latest_marker(self):
        self.assertEqual(self.upload().status_code,200)
        self.signin('lecturer','teacher')
        self.assertEqual(self.review().status_code,200)
        row=self.record();self.assertEqual((row['attendance_status'],row['absent_reason'],row['medical_status']),(0,'MC','approved'))
        self.assertEqual(self.conn.execute('SELECT updated_by FROM class_session').fetchone()[0],'Mr.Teacher')
        self.assertEqual(self.review().status_code,409)

    def test_submission_deadline_and_review_after_deadline(self):
        self.when=date.today()-timedelta(days=15)
        self.conn.execute('UPDATE class_session SET date=?',(self.when,));self.conn.execute('UPDATE student_class_session SET date=?',(self.when,))
        self.assertEqual(self.upload().status_code,400)
        self.when=date.today()-timedelta(days=14)
        self.conn.execute('UPDATE class_session SET date=?',(self.when,));self.conn.execute('UPDATE student_class_session SET date=?',(self.when,))
        self.assertEqual(self.upload().status_code,200)
        # A file submitted in time may be reviewed after the two-week period.
        original_when=self.when
        self.when-=timedelta(days=1)
        self.conn.execute('UPDATE class_session SET date=?',(self.when,));self.conn.execute('UPDATE student_class_session SET date=?,medical_submitted_at=?',(self.when,datetime.combine(original_when,datetime.min.time())))
        self.signin('hod','teacher');self.assertEqual(self.review().status_code,200)

    def test_special_requires_semester_end_and_ar_approval(self):
        self.conn.execute('UPDATE semester SET end_date=?',(date.today(),))
        self.assertEqual(self.upload('special').status_code,400)
        self.conn.execute('UPDATE semester SET end_date=?',(date.today()-timedelta(days=1),))
        self.assertEqual(self.upload('special').status_code,200)
        for role,email in [('lecturer','teacher'),('hod','head')]:
            self.signin(role,email);self.assertEqual(self.review().status_code,403)
        self.signin('ar','ar');self.assertEqual(self.review().status_code,200)
        self.assertEqual(self.record()['absent_reason'],'SMC')
        self.assertEqual(self.conn.execute('SELECT email FROM class_session').fetchone()[0],'ar')

    def test_rejection_keeps_absent_and_allows_resubmission_in_time(self):
        self.upload();self.signin('lecturer','teacher')
        self.assertEqual(self.review('rejected').status_code,200)
        self.assertIsNone(self.record()['absent_reason'])
        self.signin('student','a');self.assertEqual(self.upload().status_code,200)

    def test_present_and_changed_attendance_cannot_be_approved(self):
        self.conn.execute('UPDATE student_class_session SET attendance_status=1 WHERE email=\'a\'')
        self.assertEqual(self.upload().status_code,400)
        self.conn.execute('UPDATE student_class_session SET attendance_status=0')
        self.upload();self.conn.execute('UPDATE student_class_session SET attendance_status=1 WHERE email=\'a\'')
        self.signin('lecturer','teacher');self.assertEqual(self.review().status_code,409)
        self.assertEqual(self.record()['medical_status'],'pending')

    def test_file_validation_and_download_authorization(self):
        self.assertEqual(self.upload(content=b'<html>fake PDF</html>').status_code,400)
        self.assertEqual(self.upload(content=b'%PDF-'+b'x'*(2*1024*1024)).status_code,400)
        self.assertEqual(self.upload().status_code,200)
        from urllib.parse import urlencode
        url='/api/medical-file?'+urlencode({**self.key(),'email':'a'})
        self.assertTrue(self.client.get(url).data.startswith(b'%PDF-'))
        self.signin('student','b');self.assertEqual(self.client.get(url).status_code,404)
        self.signin('lecturer','unassigned');self.assertEqual(self.client.get(url).status_code,403)
        self.signin('student','a');self.assertEqual(self.review().status_code,403)
        self.signin('hod','head');self.assertEqual(self.upload().status_code,403)

    def test_hod_cannot_mark_or_update_unassigned_module_even_by_direct_request(self):
        from urllib.parse import urlencode
        self.signin('hod','head')
        self.assertEqual(self.client.get('/api/attendance?'+urlencode(self.key())).status_code,403)
        for mode in ('mark','update'):
            response=self.client.post('/api/attendance',json={**self.key(),'mode':mode,
                'statuses':{'ICT/2024/001':'present','ICT/2024/002':'present'},
                'originals':{'ICT/2024/001':'absent','ICT/2024/002':'absent'}},headers={'X-CSRF-Token':'token'})
            self.assertEqual(response.status_code,403)
        self.assertEqual(self.record()['attendance_status'],0)
        # Full reporting remains available even though this module is not assigned.
        self.assertEqual(self.client.get('/api/report?academic_year=26/27').status_code,200)
        self.signin('hod','teacher')
        self.assertEqual(self.client.get('/api/attendance?'+urlencode(self.key())).status_code,200)
        self.conn.execute("DELETE FROM lecturer_course WHERE email='teacher'")
        self.assertEqual(self.client.get('/api/attendance?'+urlencode(self.key())).status_code,403)

    def test_normal_medical_is_private_to_assigned_lecturer(self):
        self.upload()
        from urllib.parse import urlencode
        url='/api/medical-file?'+urlencode({**self.key(),'email':'a'})
        for role,email in [('hod','head'),('lecturer','other'),('ar','ar')]:
            self.signin(role,email)
            self.assertEqual(self.client.get('/api/medical-submissions').json['total'],0)
            self.assertEqual(self.client.get(url).status_code,403)
            self.assertEqual(self.review().status_code,403)
        self.signin('lecturer','teacher')
        self.assertEqual(self.client.get('/api/medical-submissions').json['total'],1)
        self.assertEqual(self.client.get(url).status_code,200)
        self.assertEqual(self.review().status_code,200)

    def test_pending_list_respects_assignments_and_special_restriction(self):
        self.upload('special');self.signin('lecturer','teacher')
        self.assertEqual(self.client.get('/api/medical-submissions').json['total'],0)
        self.signin('ar','ar');self.assertEqual(self.client.get('/api/medical-submissions').json['total'],1)
        self.signin('student','a');self.assertEqual(self.client.get('/api/medical-submissions').status_code,403)


if __name__=='__main__':unittest.main()
