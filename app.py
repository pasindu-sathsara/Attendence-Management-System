"""Simple Flask attendance system using MySQL and Excel module rosters."""
import json, secrets, re, io, hmac, zipfile
from pathlib import Path
from functools import wraps
from datetime import date, timedelta
from flask import Flask, request, session, jsonify, render_template, redirect, g, abort, send_from_directory, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.exceptions import HTTPException
import mysql.connector
from openpyxl import load_workbook, Workbook

ROOT = Path(__file__).resolve().parent
secret_file = ROOT / '.session-secret'
if not secret_file.exists(): secret_file.write_text(secrets.token_hex(32))
app = Flask(__name__, template_folder=str(ROOT), static_folder='static')
app.config.update(SECRET_KEY=secret_file.read_text().strip(), MAX_CONTENT_LENGTH=5*1024*1024,
                  SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax')
TABLES = {'student':'student','lecturer':'lecturer','hod':'hod','ar':'ar_user'}

BATCH_SQL = "CASE WHEN s.reg_number REGEXP '^[A-Za-z]+/[0-9]{4}/[0-9]+$' THEN CONCAT(LPAD(MOD(CAST(SUBSTRING_INDEX(SUBSTRING_INDEX(s.reg_number,'/',2),'/',-1) AS UNSIGNED)-1,100),2,'0'),'/',RIGHT(SUBSTRING_INDEX(SUBSTRING_INDEX(s.reg_number,'/',2),'/',-1),2)) ELSE NULL END"

def module_label(code):
    return code if code.endswith('(P)') or code.endswith('(T)') else code+'(T)'

def account_table():
    return 'lecturer' if session.get('role') in ('hod','lecturer') else TABLES.get(session.get('role'))


def db():
    if 'db' not in g:
        try: cfg=json.loads((ROOT/'config.json').read_text())
        except (FileNotFoundError, json.JSONDecodeError): abort(503, 'Run python setup.py to configure MySQL first.')
        g.db=mysql.connector.connect(**cfg, autocommit=False, connection_timeout=5, charset='utf8mb4')
    return g.db

def query(sql, args=(), one=False):
    cur=db().cursor(dictionary=True)
    try:
        cur.execute(sql,args)
        return cur.fetchone() if one else cur.fetchall()
    finally: cur.close()

def execute(sql,args=()):
    cur=db().cursor()
    try: cur.execute(sql,args); return cur.rowcount
    finally: cur.close()

@app.teardown_appcontext
def close_db(error):
    conn=g.pop('db',None)
    if conn: conn.close()

@app.before_request
def protect():
    session.setdefault('csrf',secrets.token_hex(32))
    if request.method in ('POST','PUT','DELETE','PATCH'):
        token=request.headers.get('X-CSRF-Token') or request.form.get('csrf','')
        if not hmac.compare_digest(str(token),session['csrf']): abort(403,'Refresh the page and try again.')

@app.after_request
def headers(response):
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='DENY'
    response.headers['Cache-Control']='no-store'
    return response

@app.errorhandler(HTTPException)
def http_error(err):
    if request.path.startswith('/api/'): return jsonify(error=err.description),err.code
    return render_template('error.html',message=err.description),err.code

@app.errorhandler(mysql.connector.Error)
def database_error(err):
    if 'db' in g: g.db.rollback()
    app.logger.error('MySQL error %s',err.errno)
    messages={1062:'This record already exists.',
              1452:'A referenced semester, course, student or staff account does not exist.',
              3819:'A value does not meet the database constraints.',
              1064:'A database query has a syntax error (1064). Please report this error to the system maintainer.',
              1054:'The database structure does not match this version. Run setup.py upgrade (1054).',
              1146:'A required database table is missing. Check the database setup (1146).'}
    connection_errors={1045,1049,2002,2003,2005,2006,2013,2055}
    if err.errno in connection_errors:
        message='Cannot connect to the database. Check config.json and that MySQL Server is running.'
        status=503
    else:
        message=messages.get(err.errno,'A database operation failed. Please report MySQL error '+str(err.errno)+'.')
        status=400 if err.errno in (1062,1452,3819) else 500
    if request.path.startswith('/api/'): return jsonify(error=message),status
    return render_template('error.html',message=message),status


def allowed(*roles):
    def decorator(fn):
        @wraps(fn)
        def wrapped(*args,**kwargs):
            if not session.get('email'):
                if request.path.startswith('/api/'): abort(401,'Please sign in.')
                return redirect('/UI.html')
            if session.get('role') not in roles: abort(403,'You do not have access to this page or action.')
            account=query('SELECT email,must_change_password FROM '+account_table()+' WHERE email=%s',(session['email'],),True)
            if not account:
                session.clear(); abort(401,'Please sign in again.')
            if account.get('must_change_password') and request.endpoint not in ('first_password','save_first_password'):
                if request.path.startswith('/api/'): abort(403,'Complete the first-login password step.')
                return redirect('/first-password.html')
            return fn(*args,**kwargs)
        return wrapped
    return decorator

def data():
    value=request.get_json(silent=True)
    if not isinstance(value,dict): abort(400,'Invalid request.')
    return value

def text(value,limit=255):
    value=str(value or '').strip()
    if not value or len(value)>limit: abort(400,'A required value is missing or too long.')
    return value

def semester_key(value):
    parts=str(value).split('|')
    if len(parts)!=3: abort(400,'Choose a semester.')
    try: return (parts[0],int(parts[1]),int(parts[2]))
    except ValueError: abort(400,'Invalid semester.')

def semester_id(row): return f"{row['academic_year']}|{row['study_year']}|{row['semester_number']}"

def course_access(value, assigned=False):
    parts=str(value).split('~')
    if len(parts)!=2: abort(400,'Select a module and academic year.')
    code,year=parts
    course=query('SELECT c.*,o.academic_year,o.teaching_end FROM course c JOIN course_offering o ON o.course_code=c.course_code WHERE c.course_code=%s AND o.academic_year=%s',(code,year),True)
    if not course: abort(404,'Module offering not found.')
    if (session['role']=='lecturer' or (assigned and session['role']=='hod')) and not query('SELECT email FROM lecturer_course WHERE email=%s AND course_code=%s AND academic_year=%s',(session['email'],code,year),True): abort(403,'This module is not assigned to you for this year.')
    return course

def roster(course):
    return query('SELECT s.email,s.reg_number FROM student s JOIN student_course e ON e.email=s.email WHERE e.course_code=%s AND e.academic_year=%s ORDER BY s.reg_number',(course['course_code'],course['academic_year']))

@app.route('/')
def home(): return redirect('/UI.html')
@app.get('/style.css')
def css(): return send_from_directory(ROOT,'style.css')
@app.get('/UI.html')
def login_page(): return render_template('UI.html', portal='student')
@app.get('/staff-login.html')
def staff_login_page(): return render_template('UI.html', portal='staff')
@app.get('/student-register.html')
def register_page(): return redirect('/UI.html')
@app.get('/staff-register.html')
def staff_register(): return redirect('/UI.html')
@app.get('/dashboard.html')
@app.get('/Dashboard.html')
@allowed('hod','lecturer','ar')
def dashboard(): return render_template('Dashboard.html')
@app.get('/student-dashboard.html')
@allowed('student')
def student_dashboard(): return render_template('student-dashboard.html')
@app.get('/manage-system.html')
@allowed('hod')
def manage(): return render_template('manage-system.html')
@app.get('/reports.html')
@allowed('hod','lecturer','ar')
def reports(): return render_template('reports.html')
@app.get('/medicals.html')
@allowed('hod','lecturer','ar')
def medicals_page(): return render_template('medicals.html')

@app.post('/api/logout')
def logout(): session.clear(); return jsonify(ok=True)

@app.post('/api/login')
def login():
    d=data(); email=text(d.get('email')).lower(); password=str(d.get('password',''))
    portal=d.get('portal')
    if portal not in ('student','staff'): abort(400,'Select the student or staff login page.')
    roles=['student'] if portal=='student' else ['lecturer','ar']
    found=None
    for role in roles:
        row=query("SELECT email,CASE WHEN NULLIF(title,'') IS NULL THEN name ELSE CONCAT(title,'.',name) END AS name,password_hash,must_change_password FROM "+TABLES[role]+' WHERE email=%s',(email,),True)
        if row:
            try: valid=check_password_hash(row['password_hash'],password)
            except (ValueError,TypeError): valid=False
            if valid:
                if role=='lecturer' and query('SELECT email FROM hod WHERE email=%s',(email,),True): role='hod'
                found=(role,row);break
    if not found: abort(401,'Incorrect email or password.')
    session.clear();session.update(email=found[1]['email'],role=found[0],name=found[1]['name'],csrf=secrets.token_hex(32))
    return jsonify(next='/first-password.html' if found[1].get('must_change_password') else '/student-dashboard.html' if found[0]=='student' else '/dashboard.html')

@app.get('/first-password.html')
@allowed('student','hod','lecturer','ar')
def first_password():
    account=query('SELECT must_change_password FROM '+account_table()+' WHERE email=%s',(session['email'],),True)
    if not account['must_change_password']: return redirect('/student-dashboard.html' if session['role']=='student' else '/dashboard.html')
    return render_template('first-password.html')

@app.post('/api/first-password')
@allowed('student','hod','lecturer','ar')
def save_first_password():
    d=data();password=str(d.get('password',''))
    if len(password)<8 or len(password)>128 or password!=d.get('confirm'):
        abort(400,'Use matching passwords of 8–128 characters. You can reuse your current password.')
    table=account_table()
    changed=execute('UPDATE '+table+' SET password_hash=%s,must_change_password=0 WHERE email=%s AND must_change_password=1',(generate_password_hash(password),session['email']))
    if not changed: abort(403,'First-login password setup is already complete. Contact the department for changes.')
    db().commit()
    return jsonify(next='/student-dashboard.html' if session['role']=='student' else '/dashboard.html')

@app.get('/api/options')
@allowed('hod','lecturer','ar')
def options():
    sql='SELECT c.*,o.academic_year FROM course c JOIN course_offering o ON o.course_code=c.course_code '
    args=()
    if session['role']=='lecturer' or (session['role']=='hod' and request.args.get('attendance')=='1'):
        sql+='WHERE EXISTS (SELECT 1 FROM lecturer_course lc WHERE lc.course_code=o.course_code AND lc.academic_year=o.academic_year AND lc.email=%s) '
        args=(session['email'],)
    courses=query(sql+'ORDER BY o.academic_year,c.study_year,c.semester_number,c.course_code',args)
    sems={semester_id(c):dict(id=semester_id(c),label=f"{c['academic_year']} · Year {c['study_year']} · Semester {c['semester_number']}") for c in courses}
    result={'semesters':list(sems.values()),'courses':[dict(code=c['course_code']+'~'+c['academic_year'],label=module_label(c['course_code'])+' · '+c['academic_year'],semester=semester_id(c)) for c in courses]}
    if session['role']=='hod':
        if request.args.get('attendance')!='1': result['semesters']=[dict(id=semester_id(c),label=f"{c['academic_year']} · Year {c['study_year']} · Semester {c['semester_number']}") for c in query('SELECT * FROM semester ORDER BY academic_year,study_year,semester_number')]
        result['lecturers']=query("SELECT email,type,CASE WHEN NULLIF(title,'') IS NULL THEN name ELSE CONCAT(title,'.',name) END AS name FROM lecturer ORDER BY name")
        result['assignments']=query("SELECT lc.email,CASE WHEN NULLIF(l.title,'') IS NULL THEN l.name ELSE CONCAT(l.title,'.',l.name) END AS name,lc.course_code,lc.academic_year FROM lecturer_course lc JOIN lecturer l ON l.email=lc.email ORDER BY l.name,lc.course_code")
    result['batches']=[r['batch'] for r in query('SELECT DISTINCT '+BATCH_SQL+' AS batch FROM student s WHERE '+BATCH_SQL+' IS NOT NULL ORDER BY batch')]
    return jsonify(result)

@app.get('/api/sessions')
@allowed('hod','lecturer','ar')
def sessions():
    course=course_access(request.args.get('course',''));code=course['course_code']
    rows=query('SELECT date,week_number,updated_by,updated_time FROM class_session WHERE course_code=%s AND academic_year=%s ORDER BY date DESC',(code,course['academic_year']))
    return jsonify([dict(id=f"{r['date']}|{r['week_number']}",label=f"{r['date']} · Week {r['week_number']}",updated_by=r['updated_by'],updated_time=str(r['updated_time'] or '')) for r in rows])

def selected_session(d,lock=False):
    course=course_access(d.get('course'),assigned=True);code=course['course_code']
    try:
        parts=str(d.get('session','')).split('|'); when=date.fromisoformat(parts[0]);week=int(parts[1])
    except (ValueError,IndexError): abort(400,'Choose a valid class session.')
    row=query('SELECT * FROM class_session WHERE course_code=%s AND date=%s AND week_number=%s'+(' FOR UPDATE' if lock else ''),(code,when,week),True)
    if not row or row['academic_year']!=course['academic_year']: abort(404,'Class session not found for this academic year.')
    return course,when,week

def attendance_status(row):
    if not row: return ''
    if row['attendance_status']==1: return 'present'
    return {'MC':'medical','SMC':'special','EXCUSE':'excused'}.get(row['absent_reason'],'absent')

def medical_rules(course,when):
    end = query('SELECT end_date FROM semester WHERE academic_year=%s AND study_year=%s AND semester_number=%s',
                (course['academic_year'], course['study_year'], course['semester_number']), True)
    return dict(normal=when <= date.today() <= when + timedelta(days=14),
                special=bool(end and end['end_date'] and date.today() > end['end_date']))


@app.get('/api/attendance')
@allowed('hod','lecturer','ar')
def attendance():
    course,when,week=selected_session(request.args)
    records=query('SELECT email,attendance_status,absent_reason FROM student_class_session WHERE course_code=%s AND date=%s AND week_number=%s',(course['course_code'],when,week))
    saved={r['email']:attendance_status(r) for r in records}
    audit=query('SELECT updated_by,updated_time FROM class_session WHERE course_code=%s AND date=%s AND week_number=%s',(course['course_code'],when,week),True)
    return jsonify(students=[dict(registrationNumber=s['reg_number'],status=saved.get(s['email'],'')) for s in roster(course)],rules=medical_rules(course,when),audit=audit)

@app.post('/api/attendance')
@allowed('hod','lecturer','ar')
def save_attendance():
    d=data();course,when,week=selected_session(d,True); students=roster(course)
    submitted=d.get('statuses'); originals=d.get('originals'); mode=d.get('mode')
    if not isinstance(submitted,dict) or mode not in ('mark','update'): abort(400,'Invalid attendance request.')
    expected={s['reg_number']:s['email'] for s in students}
    if not expected or set(submitted)!=set(expected): abort(400,'Attendance must include the entire enrolled class.')
    rows=query('SELECT * FROM student_class_session WHERE course_code=%s AND date=%s AND week_number=%s FOR UPDATE',(course['course_code'],when,week))
    old={r['email']:r for r in rows};rules=medical_rules(course,when)
    if mode=='mark' and old: abort(409,'Attendance already exists. Use Update Attendance.')
    changes=[]
    for reg,status in submitted.items():
        row=old.get(expected[reg]);current=attendance_status(row)
        if mode=='update' and (not isinstance(originals,dict) or originals.get(reg)!=current): abort(409,'Attendance changed. Reload before updating.')
        if status not in ('present','absent','medical','special','excused'): abort(400,'Mark every student.')
        if status==current: continue
        if current=='special' and session['role']!='ar': abort(403,'An AR Special Medical approval cannot be changed by academic staff.')
        if status=='excused' and (mode!='update' or current not in ('absent','medical','excused')): abort(400,'Excused is only available for a saved absence.')
        if status in ('medical','special'):
            if mode!='update' or current not in ('absent','medical','special'): abort(400,'Medical is only available for a saved absence.')
            if status=='medical' and not rules['normal']: abort(400,'Normal medical must be recorded within 14 days of the session.')
            if status=='special' and (session['role']!='ar' or not rules['special']): abort(403,'Only AR can record Special Medical after the teaching period ends.')
        if session['role']=='ar' and (mode!='update' or current not in ('absent','medical') or status not in ('medical','special')): abort(403,'AR can only approve medical for saved absences.')
        changes.append((reg,status,row,current))
    for reg,status,row,current in changes:
        reason={'medical':'MC','special':'SMC','excused':'EXCUSE'}.get(status)
        execute('INSERT INTO student_class_session (email,course_code,date,week_number,attendance_status,absent_reason) VALUES (%s,%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE attendance_status=VALUES(attendance_status),absent_reason=VALUES(absent_reason)',(expected[reg],course['course_code'],when,week,int(status=='present'),reason))
    if changes:
        execute('UPDATE class_session SET updated_by=%s,updated_time=NOW() WHERE course_code=%s AND date=%s AND week_number=%s',(session['name'],course['course_code'],when,week))
        if session['role']=='ar': execute('UPDATE class_session SET email=%s WHERE course_code=%s AND date=%s AND week_number=%s',(session['email'],course['course_code'],when,week))
    db().commit();return jsonify(ok=True)

# One row per enrolled student and module. Only sessions with saved attendance count.
ATTENDANCE_SQL = """
WITH recorded AS (
 SELECT cs.course_code,cs.academic_year,cs.date,cs.week_number
 FROM class_session cs
 WHERE {dates} AND EXISTS (
  SELECT 1 FROM student_class_session a WHERE a.course_code=cs.course_code
  AND a.date=cs.date AND a.week_number=cs.week_number)
), totals AS (
 SELECT s.email,s.reg_number AS reg,s.name,{batch} AS batch,e.course_code,e.academic_year,
 c.study_year,c.semester_number,COUNT(r.date) AS total,
 COALESCE(SUM(a.attendance_status=1),0) AS present,
 COALESCE(SUM(a.attendance_status=0 AND (a.absent_reason IS NULL OR a.absent_reason NOT IN ('MC','SMC','EXCUSE'))),0) AS absent,
 COALESCE(SUM(a.attendance_status=0 AND a.absent_reason='MC'),0) AS medical,
 COALESCE(SUM(a.attendance_status=0 AND a.absent_reason='SMC'),0) AS special,
 COALESCE(SUM(a.attendance_status=0 AND a.absent_reason='EXCUSE'),0) AS excused,
 COUNT(r.date)-COUNT(a.email) AS unmarked
 FROM student_course e JOIN student s ON s.email=e.email
 JOIN course c ON c.course_code=e.course_code
 LEFT JOIN recorded r ON r.course_code=e.course_code AND r.academic_year=e.academic_year
 LEFT JOIN student_class_session a ON a.email=e.email AND a.course_code=r.course_code
 AND a.date=r.date AND a.week_number=r.week_number
 WHERE {filters}
 GROUP BY s.email,s.reg_number,s.name,e.course_code,e.academic_year,c.study_year,c.semester_number
), scores AS (
 SELECT totals.*,present+medical+special+excused AS credited,
 100.0*(present+medical+special+excused)/NULLIF(total,0) AS raw_percentage,
 ROUND(100.0*(present+medical+special+excused)/NULLIF(total,0),2) AS percentage,
 CASE WHEN total=0 THEN 'Not assessed' WHEN (present+medical+special+excused)*100>total*80 THEN 'Eligible' ELSE 'Ineligible' END AS eligibility
 FROM totals
)
"""

@app.get('/api/my-attendance')
@allowed('student')
def own_attendance():
    student=query('SELECT name,reg_number FROM student WHERE email=%s',(session['email'],),True)
    sql=ATTENDANCE_SQL.format(batch=BATCH_SQL,dates='1=1',filters='s.email=%s')
    rows=query(sql+"SELECT *,GREATEST(0,14-total) AS remaining,ROUND(100.0*(credited+GREATEST(0,14-total))/14,2) AS maximum,CASE WHEN total>14 THEN 'Check session plan' WHEN (credited+GREATEST(0,14-total))*100<=14*80 THEN 'Cannot reach threshold' WHEN total=0 THEN 'Not assessed' WHEN credited*100>total*80 THEN 'Above threshold' ELSE 'Can recover' END AS prediction FROM scores ORDER BY academic_year DESC,study_year DESC,semester_number,course_code",(session['email'],))
    return jsonify(student=student,rows=rows)

def report_selection():
    if request.method != 'POST':
        return None
    selected = data().get('registrations')
    if selected is None:
        return None
    if not isinstance(selected, list) or not selected:
        abort(400, 'Select at least one student.')
    if any(not isinstance(reg, str) or not reg or len(reg) > 30 for reg in selected):
        abort(400, 'Invalid student selection.')
    return sorted(set(selected))


def report_scope(registrations=None):
    params = request.args
    if params.get('type', 'full') not in ('full', 'single'):
        abort(400, 'Choose Full Report or Single Report.')
    chosen = params.get('course', '')
    filters, args = [], []
    if chosen:
        course = course_access(chosen)
        filters += ['e.course_code=%s', 'e.academic_year=%s']
        args += [course['course_code'], course['academic_year']]
    elif session['role'] not in ('hod', 'ar'):
        abort(403, 'Select one of your assigned modules.')
    if session['role'] == 'lecturer' and (params.get('type') == 'single' or registrations):
        abort(403, 'Lecturers can report on all enrolled students in an assigned module.')
    for key, column in [('academic_year','e.academic_year'), ('study_year','c.study_year'),
                        ('semester_number','c.semester_number'), ('batch',BATCH_SQL)]:
        if params.get(key):
            filters.append(column + '=%s')
            args.append(params[key])
    if registrations:
        filters.append('s.reg_number IN (' + ','.join(['%s'] * len(registrations)) + ')')
        args.extend(registrations)
    if not chosen and not params.get('academic_year'):
        abort(400, 'Select an academic year.')
    dates, dateargs = [], []
    for key, operator in [('start','>='), ('end','<=')]:
        if params.get(key):
            try:
                when = date.fromisoformat(params[key])
            except ValueError:
                abort(400, 'Choose valid dates.')
            dates.append('cs.date' + operator + '%s')
            dateargs.append(when)
    if params.get('start') and params.get('end') and params['start'] > params['end']:
        abort(400, 'End date must follow start date.')
    sql = ATTENDANCE_SQL.format(batch=BATCH_SQL, dates=' AND '.join(dates) or '1=1',
                                filters=' AND '.join(filters) or '1=1')
    return sql, tuple(dateargs + args), chosen


def report_data(registrations=None):
    sql, args, chosen = report_scope(registrations)
    if request.args.get('type') == 'single':
        rows = query(sql + """SELECT reg,name,batch,SUM(present) AS present,SUM(absent) AS absent,
            SUM(total) AS total,SUM(medical) AS medical,SUM(special) AS special,SUM(excused) AS excused,
            SUM(unmarked) AS unmarked,ROUND(100.0*SUM(credited)/NULLIF(SUM(total),0),2) AS percentage,
            CASE WHEN SUM(total)=0 THEN 'Not assessed' WHEN SUM(credited)*100>SUM(total)*80
            THEN 'Eligible' ELSE 'Ineligible' END AS eligibility
            FROM scores GROUP BY email,reg,name,batch ORDER BY reg""", args)
        kind = 'single'
    elif chosen:
        rows = query(sql + 'SELECT * FROM scores ORDER BY reg', args)
        kind = 'students'
    else:
        rows = query(sql + """SELECT course_code,academic_year,study_year,semester_number,
            COUNT(*) AS students,MAX(total) AS total,SUM(present) AS present,SUM(absent) AS absent,
            SUM(medical) AS medical,SUM(special) AS special,SUM(excused) AS excused,
            SUM(unmarked) AS unmarked,ROUND(AVG(raw_percentage),2) AS percentage
            FROM scores GROUP BY course_code,academic_year,study_year,semester_number
            ORDER BY study_year,semester_number,course_code""", args)
        kind = 'modules'
    if registrations:
        found = query(sql + 'SELECT DISTINCT reg FROM scores', args)
        if {r['reg'] for r in found} != set(registrations):
            abort(400, 'Some selected students do not match these filters. Reload the student list.')
    return dict(rows=rows, kind=kind)


@app.route('/api/report', methods=['GET', 'POST'])
@allowed('hod','lecturer','ar')
def report():
    return jsonify(report_data(report_selection()))


@app.get('/api/report-students')
@allowed('hod','ar')
def report_students():
    sql, args, chosen = report_scope()
    rows = query(sql + 'SELECT DISTINCT reg,name,batch FROM scores ORDER BY reg', args)
    return jsonify(rows=rows)


def spreadsheet_table(sheet, columns, rows):
    from openpyxl.styles import Font, PatternFill, Alignment
    sheet.append([title for key,title in columns])
    for row in rows:
        sheet.append([row.get(key) for key,title in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):
                cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid',fgColor='183659')
        cell.alignment = Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[1].height = 40
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = min(40,max(16,max(len(str(c.value or '')) for c in col)+2))


STATUS_CODES = {'present':'P', 'absent':'A', 'medical':'MC', 'special':'SMC', 'excused':'EX'}


def report_registers(registrations=None):
    sql, args, chosen = report_scope(registrations)
    details = query(sql + 'SELECT * FROM scores ORDER BY course_code,academic_year,reg', args)
    if registrations and {r['reg'] for r in details} != set(registrations):
        abort(400, 'Some selected students do not match these filters. Reload the student list.')
    marks = query(sql + """SELECT t.email,t.course_code,t.academic_year,r.date,r.week_number,
        a.attendance_status,a.absent_reason FROM scores t JOIN recorded r
        ON r.course_code=t.course_code AND r.academic_year=t.academic_year
        LEFT JOIN student_class_session a ON a.email=t.email AND a.course_code=r.course_code
        AND a.date=r.date AND a.week_number=r.week_number
        ORDER BY t.course_code,r.date,r.week_number,t.reg""", args)
    staff = query("""SELECT lc.course_code,lc.academic_year,
        CASE WHEN NULLIF(l.title,'') IS NULL THEN l.name ELSE CONCAT(l.title,'.',l.name) END AS name
        FROM lecturer_course lc JOIN lecturer l ON l.email=lc.email ORDER BY l.name""")
    grouped, saved = {}, {}
    for row in details:
        key = (row['course_code'],row['academic_year'])
        grouped.setdefault(key,dict(code=key[0],year=key[1],lecturers=[],sessions=[],rows=[]))['rows'].append(row)
    for row in staff:
        key = (row['course_code'],row['academic_year'])
        if key in grouped:
            grouped[key]['lecturers'].append(row['name'])
    for row in marks:
        key = (row['course_code'],row['academic_year'])
        item = dict(date=str(row['date']),week=row['week_number'])
        if item not in grouped[key]['sessions']:
            grouped[key]['sessions'].append(item)
        saved[(key,row['email'],item['date'],item['week'])] = '—' if row['attendance_status'] is None else STATUS_CODES.get(attendance_status(row),'—')
    for key, register in grouped.items():
        for row in register['rows']:
            row['marks'] = [saved.get((key,row['email'],cs['date'],cs['week']),'—') for cs in register['sessions']]
    return list(grouped.values())


def register_workbook(registers, student=None):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    navy = PatternFill('solid',fgColor='183659')
    pale = PatternFill('solid',fgColor='EDF4FF')
    thin = Side(style='thin',color='D8E0EA')
    # Values starting with '=' must stay text, including imported names and registration numbers.
    def write(row, col, value):
        cell = sheet.cell(row,col,value)
        if isinstance(value,str):
            cell.data_type = 's'
        return cell
    def heading(row, label, columns):
        sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=columns)
        cell = write(row,1,label)
        cell.font = Font(bold=True,size=13,color='183659')
        cell.fill = pale
        sheet.row_dimensions[row].height = 25
    for register in registers:
        rows = [r for r in register['rows'] if student is None or r['reg']==student]
        if not rows:
            continue
        title = re.sub(r'[\\/*?:\[\]]','-',module_label(register['code']))[:31]
        sheet = book.create_sheet(title)
        width = 3 + len(register['sessions']) + 8
        heading(1,'Department of Computing — Rajarata University of Sri Lanka',width)
        heading(2,'Attendance Report' + (' — ' + student if student else ''),width)
        filters = ' · '.join(request.args.get(k,'') for k in ('academic_year','semester_number','batch') if request.args.get(k))
        write(3,1,filters)
        row_number = 5
        sessions = register['sessions']
        columns = 3 + len(sessions) + 8
        heading(row_number,module_label(register['code'])+' · '+register['year'],columns)
        heading(row_number+1,'Lecturer: '+(', '.join(register['lecturers']) or 'Not assigned'),columns)
        titles = ['S.No.','Registration No.','Name with Initials']
        titles += ['Week '+str(cs['week'])+'\n'+cs['date'] for cs in sessions]
        titles += ['Present','Absent','MC','Special Medical','Excused','Total Sessions','Attendance %','Eligibility']
        for col,title in enumerate(titles,1):
            cell = write(row_number+2,col,title)
            cell.font = Font(bold=True,color='FFFFFF')
            cell.fill = navy
            cell.alignment = Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[row_number+2].height = 38
        for index,record in enumerate(rows,1):
            values = [index,record['reg'],record['name']] + record['marks']
            values += [record[k] for k in ('present','absent','medical','special','excused','total','percentage','eligibility')]
            for col,value in enumerate(values,1):
                cell = write(row_number+2+index,col,value)
                cell.border = Border(bottom=thin)
                if col>3:
                    cell.alignment = Alignment(horizontal='center')
            sheet.cell(row_number+2+index,columns-1).number_format = '0.00"%"'
        row_number += len(rows)+6
        write(row_number,1,'P=Present · A=Absent · MC=Medical · SMC=Special Medical · EX=Excused · —=Not marked')
        write(row_number+1,1,'Attendance = (Present + MC + SMC + Excused) / Total Sessions x 100. Eligibility: attendance >80%.')
        sheet.column_dimensions['A'].width = 9
        sheet.column_dimensions['B'].width = 23
        sheet.column_dimensions['C'].width = 32
        for col in range(4,width+1):
            sheet.column_dimensions[get_column_letter(col)].width = 17
        sheet.freeze_panes = 'D8'
        sheet.sheet_view.showGridLines = False
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    if not book.worksheets:
        sheet = book.create_sheet('Attendance Report')
        write(1,1,'No enrolled students match these filters.')
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    return output


@app.route('/api/report.xlsx', methods=['GET','POST'])
@allowed('hod','lecturer','ar')
def export_report():
    registers = report_registers(report_selection())
    return send_file(register_workbook(registers),as_attachment=True,download_name='attendance-report.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.post('/api/student-reports.zip')
@allowed('hod','ar')
def student_reports():
    if request.args.get('type') != 'single':
        abort(400,'Choose Single Report.')
    registrations = report_selection()
    if not registrations:
        abort(400,'Select at least one student.')
    registers = report_registers(registrations)
    output = io.BytesIO()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
        for index,reg in enumerate(registrations,1):
            selected = [r for r in registers if any(s['reg']==reg for s in r['rows'])]
            workbook = register_workbook(selected,reg)
            filename = str(index)+'-'+re.sub(r'[^A-Za-z0-9_-]','_',reg)+'.xlsx'
            archive.writestr(filename,workbook.getvalue())
    output.seek(0)
    return send_file(output,as_attachment=True,download_name='student-attendance-reports.zip',mimetype='application/zip')


@app.route('/api/report-view', methods=['GET','POST'])
@allowed('hod','lecturer','ar')
def report_view():
    return jsonify(registers=report_registers(report_selection()))


@app.post('/api/semesters')
@allowed('hod')
def add_semester():
    d=data()
    try: y=int(d.get('study_year')); n=int(d.get('semester_number'))
    except (TypeError,ValueError): abort(400,'Enter valid year and semester numbers.')
    a=text(d.get('academic_year'),9)
    if not re.fullmatch(r'\d{2}/\d{2}',a) or y not in range(1,5) or n not in (1,2): abort(400,'Use academic year YY/YY, study year 1–4 and semester 1 or 2.')
    try:
        start = date.fromisoformat(d.get('start_date', ''))
        end = date.fromisoformat(d.get('end_date', ''))
    except (ValueError, TypeError):
        abort(400, 'Choose a semester start date and end date.')
    if end < start:
        abort(400, 'Semester end date must be on or after the start date.')
    execute('INSERT INTO semester (academic_year,study_year,semester_number,start_date,end_date) VALUES (%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE start_date=VALUES(start_date),end_date=VALUES(end_date)',
            (a, y, n, start, end))
    db().commit()
    return jsonify(ok=True)

@app.post('/api/courses')
@allowed('hod')
def add_course():
    d=data();key=semester_key(d.get('semester'));code=text(d.get('code'),20).upper().replace('(T)','')
    if d.get('component')=='P' and not code.endswith('(P)'): code+='(P)'
    ensure_course(code,key);db().commit();return jsonify(ok=True)

@app.post('/api/assignments')
@allowed('hod')
def assign():
    d=data();email=text(d.get('email'));course=course_access(d.get('course'))
    staff=query('SELECT type FROM lecturer WHERE email=%s',(email,),True)
    if not staff: abort(400,'Staff account not found.')
    if staff['type']=='Demonstrator' and not course['course_code'].endswith('(P)'): abort(400,'Assign demonstrators to practical components.')
    execute('INSERT INTO lecturer_course (email,course_code,academic_year) VALUES (%s,%s,%s)',(email,course['course_code'],course['academic_year']));db().commit();return jsonify(ok=True)

@app.post('/api/sessions')
@allowed('hod','lecturer','ar')
def add_session():
    d=data();course=course_access(d.get('course'),assigned=True)
    try: when=date.fromisoformat(d.get('date',''));week=int(d.get('week'))
    except (ValueError,TypeError): abort(400,'Enter a valid date and week.')
    if not 1<=week<=14: abort(400,'Week must be 1–14.')
    query('SELECT course_code FROM course_offering WHERE course_code=%s AND academic_year=%s FOR UPDATE',(course['course_code'],course['academic_year']),True)
    if query('SELECT COUNT(*) AS n FROM class_session WHERE course_code=%s AND academic_year=%s',(course['course_code'],course['academic_year']),True)['n']>=14: abort(400,'This component already has 14 sessions.')
    ar_email=session['email'] if session['role']=='ar' else None
    execute('INSERT INTO class_session (date,week_number,course_code,academic_year,email) VALUES (%s,%s,%s,%s,%s)',(when,week,course['course_code'],course['academic_year'],ar_email));db().commit();return jsonify(ok=True)

def ensure_course(code,key):
    year,study,semester=key
    existing=query('SELECT * FROM course WHERE course_code=%s',(code,),True)
    if existing and (existing['study_year']!=study or existing['semester_number']!=semester):
        abort(400,'Module '+code+' already has a different study year or semester.')
    execute('INSERT IGNORE INTO semester (academic_year,study_year,semester_number) VALUES (%s,%s,%s)',key)
    execute('INSERT IGNORE INTO course (course_code,study_year,semester_number) VALUES (%s,%s,%s)',(code,study,semester))
    execute('INSERT IGNORE INTO course_offering (course_code,academic_year) VALUES (%s,%s)',(code,year))

@app.get('/api/roster')
@allowed('hod','lecturer','ar')
def module_roster():
    return jsonify(students=[dict(registrationNumber=r['reg_number'],status='') for r in roster(course_access(request.args.get('course')))])

@app.post('/api/staff')
@allowed('hod')
def create_staff():
    d=data();email=text(d.get('email')).lower();name=text(d.get('name'),100);role=d.get('role')
    password=str(d.get('password',''))
    if role not in ('lecturer','visiting','demonstrator','ar') or '@' not in email or len(password)<8: abort(400,'Choose a role, valid email and password of at least 8 characters.')
    if any(query('SELECT email FROM '+t+' WHERE email=%s',(email,),True) for t in ('student','lecturer','ar_user')): abort(400,'Email already belongs to an account.')
    title=d.get('title')
    if title not in ('Mr','Mrs','Ms'): abort(400,'Choose a title.')
    table='ar_user' if role=='ar' else 'lecturer'
    execute('INSERT INTO '+table+' (email,name,title,password_hash) VALUES (%s,%s,%s,%s)',(email,name,title,generate_password_hash(password)))
    if table=='lecturer': execute('UPDATE lecturer SET type=%s WHERE email=%s',({'visiting':'Visiting Lecturer','demonstrator':'Demonstrator'}.get(role,'Lecturer'),email))
    db().commit();return jsonify(ok=True)

def workbook_rows(file):
    if not file or not file.filename.lower().endswith('.xlsx'): abort(400,'Choose an .xlsx workbook.')
    try:
        book=load_workbook(io.BytesIO(file.read()),read_only=True,data_only=True)
        result=[];names={};missing_names=[]
        for sheet in book:
            match=re.match(r'^\s*((?:ICT|COM|IDC))\s*(\d{4})',sheet.title,re.I)
            if not match: continue
            code=match[1].upper()+' '+match[2]+('(P)' if re.search(r'pract|\(P\)',sheet.title,re.I) else '')
            students=[];header=None
            for number,values in enumerate(sheet.iter_rows(values_only=True),1):
                if number>10000: abort(400,'A sheet exceeds 10000 rows.')
                if header is None:
                    heads=[str(v or '').strip().lower() for v in values]
                    if 'reg_no' in heads and 'name with initials' in heads:
                        header=(heads.index('reg_no'),heads.index('name with initials'))
                    continue
                if len(values)<=max(header): continue
                reg=str(values[header[0]] or '').strip().upper();name=str(values[header[1]] or '').strip()
                if not reg or '/' not in reg: continue
                prefix=sheet.title+' row '+str(number)+': '
                if len(reg)>30: abort(400,prefix+'Reg_no is too long.')
                if len(name)>100: abort(400,prefix+'Name with Initials exceeds 100 characters.')
                row=dict(reg=reg,name=name)
                if name: names[reg]=name
                else: missing_names.append((row,prefix))
                students.append(row)
            if students: result.append(dict(code=code,study=int(match[2][0]),sheet=sheet.title,students=students))
        book.close()
        for row,prefix in missing_names:
            if row['reg'] not in names: abort(400,prefix+'Name with Initials is missing for '+row['reg']+'.')
            row['name']=names[row['reg']]
        if not result: abort(400,'No module sheets with Reg_no and Name with Initials headings were found.')
        return result
    except HTTPException: raise
    except Exception: abort(400,'Unable to read the attendance workbook.')

def account_details(file, existing=None, wanted=None):
    import csv
    if not file: return {}
    result={};emails=set()
    try:
        if file.filename.lower().endswith('.csv'):
            rows=list(csv.reader(io.StringIO(file.read().decode('utf-8-sig'))))
        elif file.filename.lower().endswith('.xlsx'):
            book=load_workbook(io.BytesIO(file.read()),read_only=True,data_only=True)
            rows=list(book.active.iter_rows(values_only=True));book.close()
        else: abort(400,'Account details must be CSV or XLSX.')
        if not rows: abort(400,'Account details file is empty.')
        heads=[str(v or '').strip().lower() for v in rows[0]]
        for field in ('reg_no','email','nic'):
            if field not in heads: abort(400,'Account details row 1: missing '+field+' heading.')
        for number,values in enumerate(rows[1:],2):
            if not any(v is not None and str(v).strip() for v in values): continue
            def field(name):
                if name not in heads or heads.index(name)>=len(values): return ''
                value=values[heads.index(name)]
                if value is None: return ''
                if isinstance(value,float) and value.is_integer(): value=int(value)
                return str(value).strip()
            reg=field('reg_no').upper();email=field('email').lower();nic=field('nic')
            if not email and not nic: continue
            prefix='Account details row '+str(number)+': '
            if not reg or len(reg)>30: abort(400,prefix+'Reg_no is missing or too long.')
            if not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+',email) or len(email)>255: abort(400,prefix+'Email is invalid.')
            if not re.fullmatch(r'(?:\d{9}[vVxX]|\d{12})',nic): abort(400,prefix+'NIC must contain 12 digits, or 9 digits followed by V or X. Format NIC cells as text.')
            if reg in result: abort(400,prefix+'duplicate Reg_no '+reg)
            if email in emails: abort(400,prefix+'duplicate Email '+email)
            emails.add(email)
            password_hash = ''
            if (not existing or reg not in existing) and (wanted is None or reg in wanted):
                password_hash = generate_password_hash(nic)
            result[reg] = dict(email=email, password_hash=password_hash)
        return result
    except HTTPException: raise
    except Exception: abort(400,'Unable to read account details. Use a CSV or XLSX file with Reg_no, Email and NIC headings.')

@app.post('/api/import-preview')
@allowed('hod')
def import_preview():
    old=session.pop('import_token',None)
    if old: (ROOT/'.imports'/old).unlink(missing_ok=True)
    key=semester_key(request.form.get('semester'))
    if not re.fullmatch(r'\d{2}/\d{2}',key[0]) or key[2] not in (1,2): abort(400,'Choose an academic year and semester.')
    modules = workbook_rows(request.files.get('file'))
    registrations = sorted({r['reg'] for m in modules for r in m['students']})
    existing = {}
    for offset in range(0, len(registrations), 500):
        group = registrations[offset:offset + 500]
        saved = query('SELECT reg_number,email FROM student WHERE reg_number IN (' + ','.join(['%s'] * len(group)) + ')', tuple(group))
        existing.update({r['reg_number']: r for r in saved})
    accounts = account_details(request.files.get('accounts'), existing, set(registrations))
    rows={};members=set()
    for module in modules:
        if module['study'] not in range(1,5): abort(400,'Invalid study year in '+module['sheet'])
        for student in module['students']:
            reg=student['reg']
            if reg not in rows:
                old=existing.get(reg)
                account=accounts.get(reg)
                if old and account and old['email'].lower()!=account['email']: abort(400,'Email differs from the existing account for '+reg)
                rows[reg]=dict(reg=reg,name=student['name'],email=old['email'] if old else account['email'] if account else '',password_hash=account['password_hash'] if account and not old else '',modules=[])
            if (reg,module['code']) not in members:
                members.add((reg,module['code']));rows[reg]['modules'].append(module['code'])
    module_map={m['code']:dict(code=m['code'],study=m['study']) for m in modules}
    stage=ROOT/'.imports';stage.mkdir(exist_ok=True);token=secrets.token_hex(24)
    if session.get('import_token'): (stage/session['import_token']).unlink(missing_ok=True)
    payload=dict(rows=list(rows.values()),modules=list(module_map.values()),semester=key)
    (stage/token).write_text(json.dumps(payload));session['import_token']=token
    missing=sum(not r['email'] for r in rows.values())
    return jsonify(rows=[dict(reg=r['reg'],name=r['name'],email=r['email'] or 'Account details required',modules=', '.join(module_label(c) for c in r['modules']),study_year=max(module_map[c]['study'] for c in r['modules']),programme=student_programme(r['reg'])) for r in rows.values()],missing=missing,modules=len(module_map),sheets=len(modules))

def student_programme(reg):
    parts=re.split(r'[^A-Z0-9]+',reg.upper())
    return 'Computer Science' if any(p in ('ASP','ASB') for p in parts) else 'ICT'

def page_result(sql,args,order):
    try: page=max(1,int(request.args.get('page',1)))
    except ValueError: abort(400,'Invalid page.')
    count=query('SELECT COUNT(*) AS n FROM ('+sql+') AS results',tuple(args),True)['n']
    pages=max(1,(count+9)//10);page=min(page,pages)
    rows=query(sql+' ORDER BY '+order+' LIMIT 10 OFFSET %s',tuple(args)+((page-1)*10,))
    for row in rows:
        for key, value in row.items():
            if isinstance(value, date): row[key] = value.isoformat()
    return jsonify(rows=rows,total=count,page=page,pages=pages)

@app.get('/api/students')
@allowed('hod')
def imported_students():
    sql="""SELECT s.reg_number AS reg,s.name,s.email,"""+BATCH_SQL+""" AS batch,e.academic_year,
        MAX(c.study_year) AS study_year,
        CASE WHEN UPPER(s.reg_number) REGEXP '(^|[^A-Z0-9])(ASP|ASB)([^A-Z0-9]|$)' THEN 'Computer Science' ELSE 'ICT' END AS programme,
        GROUP_CONCAT(DISTINCT c.course_code ORDER BY c.course_code SEPARATOR ', ') AS modules
        FROM student s JOIN student_course e ON e.email=s.email JOIN course c ON c.course_code=e.course_code
        GROUP BY s.email,s.reg_number,s.name,e.academic_year"""
    filters=[];args=[]
    for key in ('academic_year','study_year','programme','batch'):
        if request.args.get(key): filters.append(key+'=%s');args.append(request.args[key])
    if request.args.get('search'):
        filters.append('(reg LIKE %s OR name LIKE %s)');args += ['%'+request.args['search']+'%']*2
    sql='SELECT * FROM ('+sql+') AS students WHERE '+(' AND '.join(filters) or '1=1')
    return page_result(sql,args,'reg')

@app.get('/api/manage-list')
@allowed('hod')
def manage_list():
    kind=request.args.get('kind'); search='%'+request.args.get('search','')+'%'
    queries={
      'modules':("SELECT c.course_code,o.academic_year,c.study_year,c.semester_number FROM course c JOIN course_offering o ON o.course_code=c.course_code WHERE c.course_code LIKE %s",'academic_year DESC,study_year,course_code'),
      'assignments':("SELECT CASE WHEN NULLIF(l.title,'') IS NULL THEN l.name ELSE CONCAT(l.title,'.',l.name) END AS name,lc.course_code,lc.academic_year FROM lecturer_course lc JOIN lecturer l ON l.email=lc.email WHERE CONCAT(l.name,' ',lc.course_code) LIKE %s",'name,course_code'),
      'staff':("SELECT * FROM (SELECT l.email,l.name,l.title,COALESCE(l.type,'Lecturer') AS staff_type,CASE WHEN h.email IS NULL THEN 'lecturer' ELSE 'hod' END AS role FROM lecturer l LEFT JOIN hod h ON h.email=l.email UNION ALL SELECT email,name,title,'AR' AS staff_type,'ar' AS role FROM ar_user) AS staff WHERE name LIKE %s",'name'),
      'semesters':("SELECT * FROM semester WHERE academic_year LIKE %s",'academic_year DESC,study_year,semester_number')}
    if kind not in queries: abort(400,'Choose a valid tab.')
    sql,order=queries[kind];return page_result(sql,[search],order)

@app.post('/api/staff-title')
@allowed('hod')
def staff_title():
    d=data();title=d.get('title');role=d.get('role')
    if title not in ('Mr','Mrs','Ms') or role not in ('lecturer','hod','ar'): abort(400,'Choose a valid title and role.')
    table='ar_user' if role=='ar' else 'lecturer'
    execute('UPDATE '+table+' SET title=%s WHERE email=%s',(title,text(d.get('email'))))
    db().commit();return jsonify(ok=True)

@app.post('/api/import-modules')
@allowed('hod')
def import_modules():
    token=session.get('import_token');path=ROOT/'.imports'/str(token)
    if not token or not path.is_file(): abort(400,'Preview a workbook first.')
    payload=json.loads(path.read_text());year,_,semester=payload['semester']
    for module in payload['modules']:
        ensure_course(module['code'],(year,module['study'],semester))
    db().commit()
    return jsonify(count=len(payload['modules']))

@app.get('/api/account-template')
@allowed('hod')
def account_template():
    import csv
    from flask import Response
    token=session.get('import_token');path=ROOT/'.imports'/str(token)
    if not token or not path.is_file(): abort(400,'Preview a workbook first.')
    out=io.StringIO();writer=csv.writer(out);writer.writerow(['Reg_no','Email','NIC'])
    for r in json.loads(path.read_text())['rows']:
        writer.writerow([r['reg'],'',''])
    return Response(out.getvalue(),mimetype='text/csv',headers={'Content-Disposition':'attachment; filename=account-details.csv'})

@app.post('/api/import-confirm')
@allowed('hod')
def import_confirm():
    token=session.get('import_token');path=ROOT/'.imports'/str(token)
    if not token or not path.is_file(): abort(400,'Preview a workbook first.')
    payload=json.loads(path.read_text());year,_,semester=payload['semester']
    if any(not r['email'] for r in payload['rows']): abort(400,'Supply account details for every new student and preview again.')
    for m in payload['modules']: ensure_course(m['code'],(year,m['study'],semester))
    studies={m['code']:m['study'] for m in payload['modules']}
    for r in payload['rows']:
        existing=query('SELECT email,reg_number FROM student WHERE email=%s OR reg_number=%s FOR UPDATE',(r['email'],r['reg']))
        if existing:
            if len(existing)!=1 or existing[0]['email'].lower()!=r['email'].lower() or existing[0]['reg_number']!=r['reg']: abort(409,'Account conflict for '+r['reg'])
        else:
            if query('SELECT email FROM lecturer WHERE email=%s UNION SELECT email FROM ar_user WHERE email=%s',(r['email'],r['email'])): abort(409,'Email belongs to a staff account.')
            execute('INSERT INTO student (email,reg_number,password_hash,name) VALUES (%s,%s,%s,%s)',(r['email'],r['reg'],r['password_hash'],r['name']))
        for code in r['modules']:
            execute('INSERT IGNORE INTO student_semester VALUES (%s,%s,%s,%s)',(r['email'],year,studies[code],semester))
            execute('INSERT IGNORE INTO student_course VALUES (%s,%s,%s)',(r['email'],code,year))
    db().commit();path.unlink();session.pop('import_token',None)
    return jsonify(ok=True,count=len(payload['rows']))


@app.post('/api/staff-type')
@allowed('hod')
def staff_type():
    d=data();kind=d.get('type')
    if kind not in ('Lecturer','Visiting Lecturer','Demonstrator'): abort(400,'Choose a staff type.')
    email=text(d.get('email'))
    if query('SELECT email FROM hod WHERE email=%s',(email,),True): abort(400,'HOD remains a lecturer.')
    if kind=='Demonstrator' and query("SELECT email FROM lecturer_course WHERE email=%s AND course_code NOT LIKE '%%(P)'",(email,),True): abort(400,'This account has theory assignments. Keep it as a lecturer.')
    execute('UPDATE lecturer SET type=%s WHERE email=%s',(kind,email));db().commit();return jsonify(ok=True)

def medical_record(values, lock=False):
    parts = str(values.get('course','')).split('~')
    if len(parts)!=2:
        abort(400,'Choose a module.')
    if session['role']=='student':
        enrolled = query('SELECT email FROM student_course WHERE email=%s AND course_code=%s AND academic_year=%s',
                         (session['email'],*parts),True)
        if not enrolled:
            abort(403,'This module is not enrolled to your account.')
        course = query('SELECT c.*,o.academic_year FROM course c JOIN course_offering o ON o.course_code=c.course_code WHERE c.course_code=%s AND o.academic_year=%s',tuple(parts),True)
    else:
        course = course_access(values.get('course'))
    if not course:
        abort(404,'Module not found.')
    try:
        when, week = str(values.get('session','')).split('|')
        when, week = date.fromisoformat(when), int(week)
    except (ValueError,TypeError):
        abort(400,'Choose a valid class session.')
    # Use the same lock order as attendance updates to avoid conflicting approvals.
    found = query('SELECT date FROM class_session WHERE course_code=%s AND academic_year=%s AND date=%s AND week_number=%s'+(' FOR UPDATE' if lock else ''),
                  (parts[0],parts[1],when,week),True)
    if not found:
        abort(404,'Class session not found.')
    email = session['email'] if session['role']=='student' else text(values.get('email'))
    row = query('SELECT * FROM student_class_session WHERE email=%s AND course_code=%s AND date=%s AND week_number=%s'+(' FOR UPDATE' if lock else ''),
                (email,parts[0],when,week),True)
    if not row:
        abort(404,'No saved attendance record was found.')
    return course,when,week,email,row


@app.post('/api/medical-upload')
@allowed('student')
def medical_upload():
    course,when,week,email,row = medical_record(request.form,True)
    kind = request.form.get('kind')
    rules = medical_rules(course,when)
    if row.get('medical_status')=='pending':
        abort(409,'A medical or excuse letter is already waiting for review.')
    status = attendance_status(row)
    if kind in ('normal','excuse'):
        if status!='absent' or not rules['normal']:
            abort(400,'Medical / Excuse Letter is only available for an absence within 14 days of the session.')
    elif kind=='special':
        if status not in ('absent','medical') or not rules['special'] or when>date.today():
            abort(400,'Special Medical is only available for an absence after the semester ends.')
    else:
        abort(400,'Choose Medical or Excuse Letter.')
    upload = request.files.get('file')
    if not upload or not upload.filename:
        abort(400,'Choose a PDF, JPG or PNG file.')
    content = upload.read(2*1024*1024+1)
    if not content or len(content)>2*1024*1024:
        abort(400,'Documents must be no larger than 2 MB.')
    types = [(b'%PDF-','application/pdf','.pdf'),(b'\xff\xd8\xff','image/jpeg','.jpg'),(b'\x89PNG\r\n\x1a\n','image/png','.png')]
    detected = next((item for item in types if content.startswith(item[0])),None)
    if not detected:
        abort(400,'Use a PDF, JPG or PNG file.')
    filename = ('excuse-letter-' if kind=='excuse' else 'medical-')+when.isoformat()+detected[2]
    execute('''UPDATE student_class_session SET medical_report=%s,medical_filename=%s,
        medical_mimetype=%s,medical_kind=%s,medical_status='pending',medical_submitted_at=NOW()
        WHERE email=%s AND course_code=%s AND date=%s AND week_number=%s''',
        (content,filename,detected[1],kind,email,course['course_code'],when,week))
    db().commit()
    return jsonify(ok=True)


def medical_access(course, kind):
    if session['role']=='student':
        return
    if kind=='special':
        if session['role']!='ar':
            abort(403,'Only AR can access Special Medical submissions.')
    elif session['role'] not in ('hod','lecturer') or not query(
            'SELECT email FROM lecturer_course WHERE email=%s AND course_code=%s AND academic_year=%s',
            (session['email'],course['course_code'],course['academic_year']),True):
        abort(403,'Only the lecturer assigned to this module can access its medical submissions.')


@app.get('/api/medical-file')
@allowed('student','hod','lecturer','ar')
def medical_file():
    course,when,week,email,row = medical_record(request.args)
    medical_access(course,row.get('medical_kind'))
    if not row.get('medical_report') or not row.get('medical_filename'):
        abort(404,'No uploaded medical file is available.')
    return send_file(io.BytesIO(row['medical_report']),as_attachment=True,
                     download_name=row['medical_filename'],mimetype=row['medical_mimetype'])


@app.get('/api/medical-submissions')
@allowed('hod','lecturer','ar')
def medical_submissions():
    filters = ["a.medical_status IS NOT NULL"]
    args = []
    if session['role']=='ar':
        filters.append("a.medical_kind='special'")
    else:
        filters += ["a.medical_kind IN ('normal','excuse')",'''EXISTS (SELECT 1 FROM lecturer_course lc
            WHERE lc.email=%s AND lc.course_code=cs.course_code AND lc.academic_year=cs.academic_year)''']
        args.append(session['email'])
    state = request.args.get('status','pending')
    if state not in ('pending','approved','rejected','all'):
        abort(400,'Choose a valid review status.')
    if state!='all':
        filters.append('a.medical_status=%s')
        args.append(state)
    sql = '''SELECT s.reg_number AS reg,s.name,a.email,a.course_code,cs.academic_year,
        a.date,a.week_number,a.medical_kind,a.medical_status,a.medical_submitted_at
        FROM student_class_session a JOIN student s ON s.email=a.email JOIN class_session cs
        ON cs.course_code=a.course_code AND cs.date=a.date AND cs.week_number=a.week_number
        WHERE '''+' AND '.join(filters)
    return page_result(sql,args,'medical_submitted_at DESC,reg')


@app.post('/api/medical-review')
@allowed('hod','lecturer','ar')
def medical_review():
    values = data()
    course,when,week,email,row = medical_record(values,True)
    medical_access(course,row.get('medical_kind'))
    if row.get('medical_status')!='pending':
        abort(409,'This submission is no longer pending. Reload the list.')
    kind = row.get('medical_kind')
    if kind not in ('normal','special','excuse'):
        abort(400,'This medical has no valid submission type.')
    if kind=='special' and session['role']!='ar':
        abort(403,'Only AR can review Special Medical.')
    decision = values.get('decision')
    if decision not in ('approved','rejected'):
        abort(400,'Choose Approve or Reject.')
    if decision=='approved':
        current = attendance_status(row)
        if (kind in ('normal','excuse') and current!='absent') or (kind=='special' and current not in ('absent','medical')):
            abort(409,'Attendance has changed. This medical can no longer be approved.')
        # Check the submission time, not the later review time.
        submitted = row['medical_submitted_at'].date()
        rules = medical_rules(course,when)
        if kind in ('normal','excuse') and not when<=submitted<=when+timedelta(days=14):
            abort(400,'This medical was not submitted within 14 days.')
        if kind=='special' and not rules['special']:
            abort(400,'The semester has not ended.')
        execute('UPDATE student_class_session SET absent_reason=%s WHERE email=%s AND course_code=%s AND date=%s AND week_number=%s',
                ({'normal':'MC','special':'SMC','excuse':'EXCUSE'}[kind],email,course['course_code'],when,week))
        execute('UPDATE class_session SET updated_by=%s,updated_time=NOW() WHERE course_code=%s AND date=%s AND week_number=%s',
                (session['name'],course['course_code'],when,week))
        if session['role']=='ar':
            execute('UPDATE class_session SET email=%s WHERE course_code=%s AND date=%s AND week_number=%s',
                    (session['email'],course['course_code'],when,week))
    execute('UPDATE student_class_session SET medical_status=%s WHERE email=%s AND course_code=%s AND date=%s AND week_number=%s',
            (decision,email,course['course_code'],when,week))
    db().commit()
    return jsonify(ok=True)


@app.get('/api/my-sessions')
@allowed('student')
def my_sessions():
    parts=str(request.args.get('course','')).split('~')
    if len(parts)!=2: abort(400,'Choose a module.')
    if not query('SELECT email FROM student_course WHERE email=%s AND course_code=%s AND academic_year=%s',(session['email'],*parts),True): abort(403,'This module is not enrolled to your account.')
    course=query('SELECT c.*,o.academic_year FROM course c JOIN course_offering o ON o.course_code=c.course_code WHERE c.course_code=%s AND o.academic_year=%s',tuple(parts),True)
    rows=query('SELECT cs.date,cs.week_number,a.attendance_status,a.absent_reason,a.medical_status,a.medical_kind,a.medical_filename FROM class_session cs LEFT JOIN student_class_session a ON a.course_code=cs.course_code AND a.date=cs.date AND a.week_number=cs.week_number AND a.email=%s WHERE cs.course_code=%s AND cs.academic_year=%s ORDER BY cs.date,cs.week_number',(session['email'],*parts))
    result=[]
    for row in rows:
        status='Not marked' if row['attendance_status'] is None else attendance_status(row)
        rules=medical_rules(course,row['date'])
        result.append(dict(date=str(row['date']),week=row['week_number'],status=status,
            medical_status=row['medical_status'],medical_kind=row['medical_kind'],has_file=bool(row['medical_filename']),
            normal=status=='absent' and rules['normal'] and row['medical_status']!='pending',
            special=status in ('absent','medical') and rules['special'] and row['date']<=date.today() and row['medical_status']!='pending'))
    return jsonify(rows=result)

if __name__=='__main__': app.run(host='127.0.0.1',port=5000,debug=False)
