"""Evaluate the report SQL against isolated SQLite fixtures; no production sample data."""
import io
import sqlite3
import unittest
from unittest.mock import patch
from werkzeug.security import generate_password_hash
from werkzeug.datastructures import FileStorage
from openpyxl import load_workbook, Workbook
import app as backend

class ReportTests(unittest.TestCase):
 def setUp(self):
  backend.app.config['TESTING']=True
  self.client=backend.app.test_client()
  with self.client.session_transaction() as s:s.update(role='hod',email='h@example.com',csrf='token')
  self.conn=sqlite3.connect(':memory:');self.conn.row_factory=sqlite3.Row
  from test_options import sql_functions
  sql_functions(self.conn)
  self.conn.executescript('''CREATE TABLE lecturer(email TEXT,name TEXT,title TEXT);
   CREATE TABLE lecturer_course(email TEXT,course_code TEXT,academic_year TEXT);
   CREATE TABLE student(email TEXT,reg_number TEXT,name TEXT);
   CREATE TABLE course(course_code TEXT,study_year INT,semester_number INT);
   CREATE TABLE course_offering(course_code TEXT,academic_year TEXT,teaching_end TEXT);
   CREATE TABLE student_course(email TEXT,course_code TEXT,academic_year TEXT);
   CREATE TABLE class_session(course_code TEXT,academic_year TEXT,date TEXT,week_number INT,updated_by TEXT);
   CREATE TABLE student_class_session(email TEXT,course_code TEXT,date TEXT,week_number INT,attendance_status INT,absent_reason TEXT);
   INSERT INTO student VALUES ('a','ICT/2024/001','Alpha'),('b','ICT/2024/002','Beta'),('c','ICT/2025/001','Gamma');
   INSERT INTO course VALUES ('ICT 1305',1,1),('COM 1301',1,1);
   INSERT INTO course_offering VALUES ('ICT 1305','26/27',NULL),('COM 1301','26/27',NULL);
   INSERT INTO student_course VALUES ('a','ICT 1305','26/27'),('b','ICT 1305','26/27'),('c','ICT 1305','26/27'),('a','COM 1301','26/27');''')
  for week in range(1,7):
   self.conn.execute('INSERT INTO class_session(course_code,academic_year,date,week_number) VALUES (?,?,?,?)',('ICT 1305','26/27',f'2099-01-{week:02}',week))
   if week==6:continue
   for email,present in [('a',week<=4),('b',True),('c',False)]:
    self.conn.execute('INSERT INTO student_class_session VALUES (?,?,?,?,?,?)',(email,'ICT 1305',f'2099-01-{week:02}',week,int(present),'MC' if email=='c' else None))
 def tearDown(self):self.conn.close()
 def query(self,sql,args=(),one=False):
  if sql.startswith('SELECT email,must_change_password FROM') or sql.startswith('SELECT email FROM hod'):return {'email':'h@example.com'}
  if sql.startswith('SELECT email FROM lecturer WHERE'):return {'email':'l@example.com'}
  if sql.startswith('SELECT email FROM lecturer_course'):return None
  cursor=self.conn.execute(sql.replace('RIGHT(', 'RIGHT_STR(').replace('%s','?'),args)
  if one:
   row=cursor.fetchone();return dict(row) if row else None
  return [dict(r) for r in cursor.fetchall()]
 def get(self,path):
  with patch.object(backend,'query',side_effect=self.query):return self.client.get(path)
 def test_individual_counts_eligibility_and_future_saved_sessions(self):
  r=self.get('/api/report?course=ICT+1305~26/27&batch=23/24')
  self.assertEqual(r.status_code,200)
  a,b=r.json['rows'];self.assertEqual(a['total'],5);self.assertEqual(a['present'],4)
  self.assertEqual(a['absent'],1);self.assertEqual(a['percentage'],80);self.assertEqual(a['eligibility'],'Ineligible')
  self.assertEqual(b['percentage'],100);self.assertEqual(b['eligibility'],'Eligible')
 def test_hod_batch_average_is_average_of_students(self):
  r=self.get('/api/report?academic_year=26/27&batch=23/24')
  row=next(r for r in r.json['rows'] if r['course_code']=='ICT 1305')
  self.assertEqual(row['students'],2);self.assertEqual(row['percentage'],90)
  r=self.get('/api/report?academic_year=26/27&batch=24/25')
  row=next(r for r in r.json['rows'] if r['course_code']=='ICT 1305')
  self.assertEqual(row['percentage'],100);self.assertEqual(row['students'],1)
 def test_empty_unenrolled_module_is_not_in_full_report(self):
  self.conn.execute("INSERT INTO course VALUES ('EMPTY',1,1)")
  self.conn.execute("INSERT INTO course_offering VALUES ('EMPTY','26/27',NULL)")
  r=self.get('/api/report?academic_year=26/27')
  self.assertFalse(any(row['course_code']=='EMPTY' for row in r.json['rows']))
 def test_no_sessions_is_not_assessed(self):
  row=self.get('/api/report?course=COM+1301~26/27').json['rows'][0]
  self.assertEqual(row['total'],0);self.assertIsNone(row['percentage']);self.assertEqual(row['eligibility'],'Not assessed')
 def test_mc_counts_towards_attendance_and_is_separate_from_absent(self):
  row=self.get('/api/report?course=ICT+1305~26/27&batch=24/25').json['rows'][0]
  self.assertEqual(row['absent'],0);self.assertEqual(row['present'],0);self.assertEqual(row['medical'],5)
  self.assertEqual(row['percentage'],100);self.assertEqual(row['eligibility'],'Eligible')
 def test_dates_limit_denominator(self):
  row=self.get('/api/report?course=ICT+1305~26/27&batch=23/24&start=2099-01-05&end=2099-01-05').json['rows'][0]
  self.assertEqual(row['total'],1);self.assertEqual(row['present'],0)
  self.assertEqual(self.get('/api/report?course=ICT+1305~26/27&start=2099-02-01&end=2099-01-01').status_code,400)
 def test_lecturer_cannot_request_all_or_unassigned_exports(self):
  with self.client.session_transaction() as s:s['role']='lecturer'
  for path in ['/api/report?academic_year=26/27','/api/report.xlsx?academic_year=26/27','/api/report?course=ICT+1305~26/27','/api/report.xlsx?course=ICT+1305~26/27']:
   self.assertEqual(self.get(path).status_code,403)
 def test_excel_matches_sql_and_has_expected_columns(self):
  response=self.get('/api/report.xlsx?course=ICT+1305~26/27&batch=23/24')
  self.assertEqual(response.status_code,200)
  book=load_workbook(io.BytesIO(response.data));rows=list(book.active.values)
  self.assertEqual(book.sheetnames,['ICT 1305(T)'])
  self.assertEqual(rows[6][:3],('S.No.','Registration No.','Name with Initials'))
  self.assertEqual(rows[7][:3],(1,'ICT/2024/001','Alpha'))
  self.assertEqual(rows[7][3:8],('P','P','P','P','A'))
  self.assertEqual(rows[7][-8:],(4,1,0,0,0,5,80,'Ineligible'))
  self.assertEqual(rows[8][1],'ICT/2024/002')
 def test_unmarked_not_invented_as_present_or_absent(self):
  self.conn.execute("DELETE FROM student_class_session WHERE email='a' AND week_number=5")
  row=self.get('/api/report?course=ICT+1305~26/27&batch=23/24').json['rows'][0]
  self.assertEqual(row['total'],5);self.assertEqual(row['present'],4);self.assertEqual(row['absent'],0);self.assertEqual(row['unmarked'],1)

class ChangesTests(unittest.TestCase):
 def setUp(self):
  self.client=backend.app.test_client()
  with self.client.session_transaction() as s:s.update(role='hod',email='h@example.com',csrf='token')
 def post(self,url,payload):return self.client.post(url,json=payload,headers={'X-CSRF-Token':'token'})
 def test_all_staff_roles_create_sessions_without_ar_lookup(self):
  for role in ('hod','lecturer','ar'):
   with self.client.session_transaction() as s:s['role']=role
   with patch.object(backend,'query',return_value={'email':'h@example.com','n':0}) as q,patch.object(backend,'course_access',return_value={'course_code':'ICT 1305','academic_year':'26/27'}),patch.object(backend,'execute') as execute,patch.object(backend,'db'):
    r=self.post('/api/sessions',{'course':'ICT 1305~26/27','date':'2026-10-24','week':1})
    self.assertEqual(r.status_code,200);self.assertEqual(q.call_count,3)
    self.assertEqual(execute.call_args.args[1][-1],'h@example.com' if role=='ar' else None)
 def test_student_portal_does_not_authenticate_staff(self):
  with patch.object(backend,'query',return_value=None) as query:
   self.assertEqual(self.post('/api/login',{'portal':'student','email':'staff@example.com','password':'password'}).status_code,401)
   self.assertEqual(query.call_count,1);self.assertIn('FROM student',query.call_args.args[0])
 def test_numeric_nic_xlsx_and_staff_title(self):
  book=Workbook();book.active.append(['Reg_no','Email','NIC','Batch']);book.active.append(['ASP/2024/001','s@example.com',200012345678,'23/24']);out=io.BytesIO();book.save(out);out.seek(0)
  rows=backend.account_details(FileStorage(stream=out,filename='details.xlsx'))
  self.assertNotIn('batch',rows['ASP/2024/001'])
  with patch.object(backend,'query',side_effect=[{'email':'h@example.com'},None,None,None]),patch.object(backend,'execute') as execute,patch.object(backend,'db'):
   response=self.post('/api/staff',{'role':'lecturer','email':'v@example.com','name':'Name','title':'Mr','password':'password123'})
   self.assertEqual(response.status_code,200);self.assertEqual(execute.call_args_list[0].args[1][2],'Mr')
 def test_import_error_identifies_row_and_field(self):
  from werkzeug.exceptions import BadRequest
  f=FileStorage(stream=io.BytesIO(b'Reg_no,Email,NIC\nA,s@example.com,2.00012E+11\n'),filename='details.csv')
  with self.assertRaises(BadRequest) as caught:backend.account_details(f)
  self.assertIn('row 2: NIC',caught.exception.description)
 def test_batch_update_only_hod(self):
  for role in ('student','lecturer','ar'):
   with self.client.session_transaction() as s:s['role']=role
   self.assertEqual(self.post('/api/student-batch',{'batch':'23/24','registrations':['A']}).status_code,404)

if __name__=='__main__':unittest.main()
